"""SS07 两张数学表设计规格生成（目标指标 spec，Word）—— 只用 SS07 自有指标。

最终设计（命中率解冻；两张都砍 15x；底层 paytable 不变）：
  v3 (95%) 低倍·高命中·稳健版：重心 2x/5x，命中率高、中奖频繁、单奖小。
  v2 (97%) 中倍·大奖版：重心 5x/10x，命中率较低（同 RTP 下高倍必然少命中）、单奖大。
  两张都把 x15 砍到很低（现状 .045/.073 → .01~.03）。
RTP = 命中率 × 单次赔付；单次赔付 = 底层 paytable(不变) × E[Applied Multiplier]；命中率反解 RTP。
"""
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[2]
MTD = ROOT / "math_table_design"
SS07 = MTD / "SS07 Math Tables Summary Statistics .xlsx"
OUT = MTD / "SS07数学表设计规格_v1.0.docx"
CN = "PingFang SC"; ACCENT = RGBColor(0x09, 0x69, 0xDA); MUTED = RGBColor(0x65, 0x6D, 0x76)
MT = [1, 2, 3, 5, 10, 15]
RTP_CUR = {"Normal": 0.964628, "Extra": 0.964757}
HIT_CUR = {"Normal": 0.120032, "Extra": 0.120678}
BET = {"Normal": 10, "Extra": 15}
RTP_T = {"v2": 0.97, "v3": 0.95}
# 目标倍率分布（形状驱动；v3=2x/5x重心 / v2=5x/10x重心；两者都砍15x）
SHAPE = {
    "v3": {"Normal": [0.260, 0.320, 0.160, 0.200, 0.050, 0.010],
           "Extra":  [0.000, 0.460, 0.220, 0.240, 0.070, 0.010]},
    "v2": {"Normal": [0.150, 0.180, 0.150, 0.280, 0.220, 0.020],
           "Extra":  [0.000, 0.220, 0.170, 0.320, 0.260, 0.030]},
}

s = pd.read_excel(SS07, sheet_name="Summary", header=None); s[0] = s[0].ffill()
def row(k1, k2):
    for _, r in s.iterrows():
        if str(r[0]) == k1 and (k2 is None or (isinstance(r[1], str) and k2 in r[1])):
            return r[2], r[3]
    return (None, None)
cur = {"Normal": [], "Extra": []}
for _, r in s[s[0] == "Applied Multiplier"].iterrows():
    if isinstance(r[1], str) and r[1].startswith("x"):
        cur["Normal"].append(float(r[2])); cur["Extra"].append(float(r[3]))

def Em(w): return sum(wi * mi for wi, mi in zip(w, MT))
basepay = {lv: (RTP_CUR[lv] / HIT_CUR[lv]) / Em(cur[lv]) for lv in cur}
prof = {}
for lv in ("Normal", "Extra"):
    prof[lv] = {"cur": {"hit": HIT_CUR[lv], "dist": cur[lv], "Em": Em(cur[lv]),
                        "payhit": RTP_CUR[lv] / HIT_CUR[lv]}}
    for tab in ("v2", "v3"):
        w = SHAPE[tab][lv]; ssum = sum(w); w = [x / ssum for x in w]; E = Em(w)
        hit = RTP_T[tab] / (basepay[lv] * E)
        prof[lv][tab] = {"hit": hit, "dist": w, "Em": E, "payhit": basepay[lv] * E}

# ---------- Word ----------
doc = Document()
nm = doc.styles["Normal"]; nm.font.name = "Calibri"; nm.font.size = Pt(10.5)
nm.element.rPr.rFonts.set(qn("w:eastAsia"), CN)
def _cn(r, size=None, bold=None, color=None):
    r.font.name = "Calibri"; r._element.rPr.rFonts.set(qn("w:eastAsia"), CN)
    if size: r.font.size = Pt(size)
    if bold is not None: r.font.bold = bold
    if color is not None: r.font.color.rgb = color
def H(t, lv=1):
    p = doc.add_paragraph(); _cn(p.add_run(t), 17 if lv == 1 else 13, True, None if lv == 1 else ACCENT)
    if lv == 1:
        pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
        for kk, vv in [("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", "1F2328")]: bot.set(qn(kk), vv)
        b.append(bot); pPr.append(b)
def P(t, size=10.5, bold=False, color=None, italic=False):
    p = doc.add_paragraph(); r = p.add_run(t); _cn(r, size, bold, color); r.font.italic = italic; return p
def shade(c, hexc):
    tcPr = c._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexc); tcPr.append(sh)
def T(headers, rows, hi=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; _cn(c.paragraphs[0].add_run(h), 9.5, True); shade(c, "F0F3F6")
    for rr in rows:
        cs = t.add_row().cells
        for i, v in enumerate(rr):
            cs[i].text = ""; _cn(cs[i].paragraphs[0].add_run(str(v)), 9.5)
            if hi and i in hi: shade(cs[i], "FFF3B0")
    doc.add_paragraph(); return t

P("SS07 数学表设计规格 · v1.0", size=18, bold=True)
for m in ["适用游戏：SS07 · 指标全用 SS07 自有 summary（不引入 base/free/scatter）",
          "SS07_Saitekika_v3（RTP 95.0%，低倍·高命中·稳健版）｜ SS07_Saitekika_v2（RTP 97.0%，中倍·大奖版）",
          "Normal(bet10) / Extra(bet15) · 命中率解冻；两张都砍 15x · 起草 2026-08-14 · draft"]:
    P(m, size=9, color=MUTED)

H("1. 两张表的特性（核心）")
P("两张都削掉 15x 以上的极端尾部；区别在中低倍率重心与命中率：", size=10)
T(["维度", "SS07_Saitekika_v3 (95%)  低倍·高命中", "SS07_Saitekika_v2 (97%)  中倍·大奖"], [
    ["RTP", "95.0%", "97.0%"],
    ["倍率重心", "2x / 5x（低倍）", "5x / 10x（中倍）"],
    ["命中率 Hit%", f"N {prof['Normal']['v3']['hit']*100:.2f}% / E {prof['Extra']['v3']['hit']*100:.2f}%（高，中奖频繁）",
     f"N {prof['Normal']['v2']['hit']*100:.2f}% / E {prof['Extra']['v2']['hit']*100:.2f}%（较低，少而大）"],
    ["单次赔付(×bet)", f"N {prof['Normal']['v3']['payhit']:.2f} / E {prof['Extra']['v3']['payhit']:.2f}（小）",
     f"N {prof['Normal']['v2']['payhit']:.2f} / E {prof['Extra']['v2']['payhit']:.2f}（大）"],
    ["E[Applied Mult]", f"N {prof['Normal']['v3']['Em']:.2f} / E {prof['Extra']['v3']['Em']:.2f}",
     f"N {prof['Normal']['v2']['Em']:.2f} / E {prof['Extra']['v2']['Em']:.2f}"],
    ["x15（都砍）", f"{SHAPE['v3']['Normal'][5]:.3f} / {SHAPE['v3']['Extra'][5]:.3f}",
     f"{SHAPE['v2']['Normal'][5]:.3f} / {SHAPE['v2']['Extra'][5]:.3f}"],
    ["波动 SD", "较低（低倍高命中）", "较高（中倍少命中）"],
    ["体感", "频繁小奖、稳健", "少而大、爽点集中在 5x/10x"],
], hi=[1, 2])
P(f"现状 x15 权重 Normal {cur['Normal'][5]:.3f} / Extra {cur['Extra'][5]:.3f} —— 两表均大幅下调。"
  f"底层 paytable 不变；命中率由 RTP 反解（RTP = 命中率 × 单次赔付）。", size=9.5, color=MUTED)

H("2. 目标总表（SS07 指标行）")
T(["指标", "现状 N / E", "v3 (95%) N / E", "v2 (97%) N / E"], [
    ["Bet", "10 / 15", "10 / 15", "10 / 15"],
    ["RTP Mean", f"{RTP_CUR['Normal']*100:.2f}% / {RTP_CUR['Extra']*100:.2f}%", "95.00% / 95.00%", "97.00% / 97.00%"],
    ["Hit% Base Game", f"{HIT_CUR['Normal']*100:.2f}% / {HIT_CUR['Extra']*100:.2f}%",
     f"{prof['Normal']['v3']['hit']*100:.2f}% / {prof['Extra']['v3']['hit']*100:.2f}%",
     f"{prof['Normal']['v2']['hit']*100:.2f}% / {prof['Extra']['v2']['hit']*100:.2f}%"],
    ["E[Applied Mult]", f"{prof['Normal']['cur']['Em']:.2f} / {prof['Extra']['cur']['Em']:.2f}",
     f"{prof['Normal']['v3']['Em']:.2f} / {prof['Extra']['v3']['Em']:.2f}",
     f"{prof['Normal']['v2']['Em']:.2f} / {prof['Extra']['v2']['Em']:.2f}"],
    ["单次赔付(×bet)", f"{prof['Normal']['cur']['payhit']:.2f} / {prof['Extra']['cur']['payhit']:.2f}",
     f"{prof['Normal']['v3']['payhit']:.2f} / {prof['Extra']['v3']['payhit']:.2f}",
     f"{prof['Normal']['v2']['payhit']:.2f} / {prof['Extra']['v2']['payhit']:.2f}"],
], hi=[2, 3])

H("3. Applied Multiplier 目标分布")
for lv in ("Normal", "Extra"):
    P(f"3.{1 if lv=='Normal' else 2} {lv}（bet {BET[lv]}）", bold=True, color=ACCENT)
    rows = []
    for tab, label in [("cur", "现状"), ("v3", "v3 (95%) 低倍"), ("v2", "v2 (97%) 中倍")]:
        w = prof[lv][tab]["dist"]; rows.append([label] + [f"{x:.3f}" for x in w] + [f"{Em(w):.3f}"])
    T(["分布"] + [f"x{m}" for m in MT] + ["E[mult]"], rows, hi=[6])
P("v3 重心压到 x2/x5、x15≈0.01；v2 重心抬到 x5/x10、x15≈0.02-0.03。两表 x15 均远低于现状。Extra 保持“无 x1”。",
  size=9.5, color=MUTED)

H("4. 开放项")
T(["编号", "问题"], [
    ["Q1", "两表倍率形状/命中率（v3 N~12.6% 2x/5x、v2 N~8% 5x/10x）是否合适，可继续微调重心与命中率"],
    ["Q2", "Max Multiplier / 头奖：两表都砍 15x 后极端头奖下降，v2>v3；具体上限由 reel 仿真定"],
    ["Q3", "底层 paytable 保持不变（现假设）；Normal/Extra 是否共用同一 reel（Extra 仅去 x1）"],
    ["Q4", "交付后数学团队据本 spec 搭 reel/paytable 并仿真回验：命中目标 RTP±0.1%、命中率、倍率分布达标"],
])
P("SS07 数学表设计规格 v1.0 · 2026-08-14 · draft", size=8.5, color=MUTED)
doc.save(str(OUT))
print("saved:", OUT)
